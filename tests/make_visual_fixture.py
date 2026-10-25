"""Generate a manual PowerPoint QA fixture; pass an output directory."""
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.korean_pptx import create_korean_card, set_native_bullet, apply_korean_formatting, set_korean_font_and_lang
from scripts.fix_korean_pptx import process_pptx

def make_fixture(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    title = sl.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12), Inches(.7))
    title.text_frame.text = "한국어 서식 검증  |  줄바꿈 · 불릿 · 표 · 그룹"
    title.text_frame.paragraphs[0].font.size = Pt(28)
    card = create_korean_card(sl, Inches(.5), Inches(1.3), Inches(6), Inches(4.6))
    tf = card.text_frame
    tf.text = "편집 가능한 단일 카드"
    tf.paragraphs[0].font.size = Pt(22)
    tf.paragraphs[0].font.bold = True
    p = tf.add_paragraph()
    p.text = "한국어 문장은 단어 단위로 줄바꿈합니다. 긴 문장에서 불릿 둘째 줄이 본문 시작 위치에 맞는지 확인합니다."
    p.font.size = Pt(18)
    set_native_bullet(p)
    p = tf.add_paragraph()
    p.text = "두 번째 항목은 같은 도형 안에서 문단으로 편집됩니다."
    p.font.size = Pt(18)
    set_native_bullet(p)
    p = tf.add_paragraph()
    p.font.size = Pt(16)
    r=p.add_run(); r.text="한국어 + "
    r=p.add_run(); r.text="English 2026"; r.font.name="Arial"
    r.font.language_id=1033
    table = sl.shapes.add_table(3,2,Inches(7),Inches(1.3),Inches(5.7),Inches(2.7)).table
    for i,row in enumerate((("검증 항목","기대 결과"),("표 내부 한국어","상속 유지"),("줄바꿈·불릿","넘침 없음"))):
        for j,t in enumerate(row):
            table.cell(i,j).text=t
            table.cell(i,j).text_frame.paragraphs[0].font.size=Pt(18)
    group=sl.shapes.add_group_shape()
    for x,txt in ((7,"그룹 도형 A"),(10,"그룹 도형 B")):
        shape=group.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(4.5),Inches(2.7),Inches(1))
        shape.text_frame.text=txt
        shape.text_frame.paragraphs[0].font.size=Pt(18)
    footer=sl.shapes.add_textbox(Inches(.5), Inches(6.4), Inches(12), Inches(.5))
    footer.text_frame.text="필드: "
    p=footer.text_frame.paragraphs[0]
    p._p.append(parse_xml('<a:fld '+nsdecls('a')+' id="{00000000-0000-0000-0000-000000000001}" type="datetime1"><a:rPr lang="en-US"/><a:t>2026년 10월 25일</a:t></a:fld>'))
    p.font.size=Pt(16)
    apply_korean_formatting(sl,font_name="맑은 고딕")

    sl=prs.slides.add_slide(prs.slide_layouts[0])
    sl.shapes.title.text="마스터 상속을 유지하는 제목"
    sl.placeholders[1].text="이 슬라이드는 폰트를 직접 지정하지 않고 기본 보정만 적용했습니다."
    apply_korean_formatting(sl)

    sl=prs.slides.add_slide(prs.slide_layouts[6])
    title=sl.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12), Inches(.7))
    title.text_frame.text="문단 기본 폰트 변경 후에도 run 상속 유지"
    title.text_frame.paragraphs[0].font.size=Pt(26)
    for x,font,label in ((.5,"맑은 고딕","맑은 고딕"),(7,"바탕","바탕으로 상위 폰트 변경")):
        shape=create_korean_card(sl,Inches(x), Inches(1.7), Inches(5.8), Inches(3.8))
        shape.text_frame.text=label+"\n같은 한국어 본문입니다.\n개별 run 폰트는 비어 있습니다."
        for p in shape.text_frame.paragraphs:
            p.font.size=Pt(24)
        apply_korean_formatting(shape,font_name="맑은 고딕")
        for p in shape.text_frame.paragraphs:
            set_korean_font_and_lang(p,font)
            assert all(r.font.name is None for r in p.runs)
    apply_korean_formatting(sl)
    raw=folder/"visual-before.pptx"
    output=folder/"visual-verified.pptx"
    # Seed a regression so the displayed output exercises CLI XML mutation too.
    card.text_frame.word_wrap = False
    card.text_frame.paragraphs[1]._p.get_or_add_pPr().set("eaLnBrk", "1")
    card.text_frame.paragraphs[1].runs[0].font.language_id = 1033
    prs.save(raw)
    assert process_pptx(raw,output_path=output)==0
    assert len(Presentation(output).slides)==3
    return output

if __name__=="__main__":
    print(make_fixture(sys.argv[1]))
