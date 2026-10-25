# -*- coding: utf-8 -*-
import contextlib
import io
import unittest
import os
import sys
import tempfile
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR

SKILL_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if SKILL_ROOT not in sys.path:
    sys.path.insert(0, SKILL_ROOT)

from scripts.korean_pptx import (
    create_korean_card,
    create_korean_shape,
    set_native_bullet,
    set_korean_font_and_lang,
    set_word_break_prevention,
    apply_korean_formatting
)
from scripts.fix_korean_pptx import (
    patch_xml_content,
    verify_xml_content,
    process_pptx,
    get_slide_file_map,
    resolve_ea_typeface,
    qualify,
    A_NS
)


class TestKoreanPptxHelper(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def test_create_korean_card(self):
        card = create_korean_card(self.slide, Inches(1), Inches(1), Inches(4), Inches(2))
        self.assertTrue(card.has_text_frame)
        self.assertTrue(card.text_frame.word_wrap)
        self.assertEqual(card.text_frame.margin_left, Inches(0.3))
        self.assertEqual(card.text_frame.vertical_anchor, MSO_ANCHOR.TOP)
        self.assertEqual(card.auto_shape_type, MSO_SHAPE.RECTANGLE)

    def test_create_korean_card_oval(self):
        oval = create_korean_card(
            self.slide, Inches(2), Inches(2), Inches(2), Inches(2),
            shape_type=MSO_SHAPE.OVAL
        )
        self.assertTrue(oval.has_text_frame)
        self.assertEqual(oval.auto_shape_type, MSO_SHAPE.OVAL)
        self.assertEqual(oval.text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)

    def test_create_korean_shape_alias(self):
        badge = create_korean_shape(
            self.slide, Inches(1), Inches(1), Inches(2), Inches(0.5),
            shape_type=MSO_SHAPE.ROUNDED_RECTANGLE,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
            margin_inch=0.1
        )
        self.assertTrue(badge.has_text_frame)
        self.assertEqual(badge.auto_shape_type, MSO_SHAPE.ROUNDED_RECTANGLE)
        self.assertEqual(badge.text_frame.vertical_anchor, MSO_ANCHOR.MIDDLE)
        self.assertEqual(badge.text_frame.margin_left, Inches(0.1))

    def test_set_native_bullet(self):
        tb = self.slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        p = tb.text_frame.paragraphs[0]
        p.text = "불릿 항목입니다."
        set_native_bullet(p, char="➔")

        pPr = p._p.get_or_add_pPr()
        tags = [c.tag.split('}')[-1] for c in pPr]
        self.assertIn("buChar", tags)
        if "defRPr" in tags:
            self.assertLess(tags.index("buChar"), tags.index("defRPr"))

    def test_set_korean_font_and_lang_paragraph(self):
        tb = self.slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        p = tb.text_frame.paragraphs[0]
        p.text = "문단 레벨 서식 상속 테스트"
        set_korean_font_and_lang(p, "Pretendard")

        defRPr = p._p.find('.//' + qualify(A_NS, 'defRPr'))
        self.assertIsNotNone(defRPr)
        ea = defRPr.find(qualify(A_NS, 'ea'))
        self.assertIsNotNone(ea)
        self.assertEqual(ea.get("typeface"), "Pretendard")

    def test_apply_korean_formatting_table(self):
        tbl_shape = self.slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        cell = tbl_shape.table.cell(0, 0)
        cell.text = "표 내부 한국어 텍스트"

        apply_korean_formatting(self.slide, font_name="Pretendard")

        p = cell.text_frame.paragraphs[0]
        self.assertEqual(p._p.get_or_add_pPr().get("eaLnBrk"), "0")
        self.assertEqual(p._p.get_or_add_pPr().get("latinLnBrk"), "0")
        ea = p.runs[0]._r.find(".//" + qualify(A_NS, "ea"))
        self.assertIsNotNone(ea)
        self.assertEqual(ea.get("typeface"), "Pretendard")

    def test_apply_korean_formatting_presentation(self):
        tb = self.slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        tb.text_frame.text = "전체 덱 적용 테스트"

        apply_korean_formatting(self.prs, font_name="Pretendard")
        p = tb.text_frame.paragraphs[0]
        self.assertEqual(p._p.get_or_add_pPr().get("eaLnBrk"), "0")


class TestFixKoreanPptxCli(unittest.TestCase):
    def test_patch_xml_missing_pPr(self):
        xml_str = """<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:txBody>
    <a:bodyPr/>
    <a:p>
      <a:r>
        <a:t>pPr이 없는 문단</a:t>
      </a:r>
    </a:p>
  </p:txBody>
</p:sp>"""
        mod, new_xml, stats = patch_xml_content(xml_str.encode("utf-8"))
        self.assertTrue(mod)
        root = ET.fromstring(new_xml)
        p = root.find('.//' + qualify(A_NS, 'p'))
        pPr = p.find(qualify(A_NS, 'pPr'))
        self.assertIsNotNone(pPr)
        self.assertEqual(pPr.get("eaLnBrk"), "0")
        self.assertEqual(pPr.get("latinLnBrk"), "0")

    def test_patch_xml_schema_order_with_hyperlink(self):
        xml_str = """<p:sp xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:txBody>
    <a:bodyPr/>
    <a:p>
      <a:pPr/>
      <a:r>
        <a:rPr>
          <a:latin typeface="Pretendard"/>
          <a:hlinkClick r:id="rId2"/>
        </a:rPr>
        <a:t>링크 텍스트</a:t>
      </a:r>
    </a:p>
  </p:txBody>
</p:sp>"""
        mod, new_xml, stats = patch_xml_content(xml_str.encode("utf-8"))
        self.assertTrue(mod)
        root = ET.fromstring(new_xml)
        rPr = root.find('.//' + qualify(A_NS, 'rPr'))
        tags = [c.tag.split('}')[-1] for c in rPr]
        self.assertEqual(tags, ['latin', 'ea', 'hlinkClick'])

    def test_full_pipeline_process_pptx(self):
        temp_dir = tempfile.mkdtemp()
        pptx_path = os.path.join(temp_dir, 'pipeline_test.pptx')

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        tb.text_frame.text = "파이프라인 테스트"
        prs.save(pptx_path)

        # 1. 보정 전 verify는 미흡 항목 발견 (코드 2)
        code_before = process_pptx(pptx_path, verify_only=True)
        self.assertEqual(code_before, 2)

        # 2. 보정 실행 (코드 0)
        code_fix = process_pptx(pptx_path)
        self.assertEqual(code_fix, 0)

        # 3. 보정 후 verify는 완벽 적용 (코드 0)
        code_after = process_pptx(pptx_path, verify_only=True)
        self.assertEqual(code_after, 0)


class TestRegressions(unittest.TestCase):
    """검토에서 재현된 결함 1~6번 회귀 테스트"""

    def _make_pptx(self):
        temp_dir = tempfile.mkdtemp()
        path = os.path.join(temp_dir, "regression.pptx")
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        tb.text_frame.text = "회귀 테스트"
        prs.save(path)
        return path

    # 1. cp949 콘솔에서 출력 문자로 크래시하지 않아야 함
    def test_cli_output_encodable_in_cp949(self):
        path = self._make_pptx()
        for verify_only in (True, False, True):
            buf = io.TextIOWrapper(io.BytesIO(), encoding="cp949", errors="strict")
            with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
                code = process_pptx(path, verify_only=verify_only)
            buf.flush()
            self.assertIn(code, (0, 2))

    # 2. 네임스페이스 접두사와 mc:Ignorable/Requires 참조가 보존되어야 함
    def test_patch_preserves_namespace_prefixes(self):
        xml_str = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" xmlns:v="urn:schemas-microsoft-com:vml" mc:Ignorable="p14 v"><p:cSld><p:spTree><p:sp><p:txBody><a:bodyPr/><a:p><a:r><a:t>한글</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld><mc:AlternateContent><mc:Choice Requires="p14"><p:transition p14:dur="10"/></mc:Choice><mc:Fallback><p:transition/></mc:Fallback></mc:AlternateContent></p:sld>"""
        mod, new_xml, _ = patch_xml_content(xml_str.encode("utf-8"))
        self.assertTrue(mod)
        out = new_xml.decode("utf-8")
        self.assertIn('xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main"', out)
        self.assertIn('xmlns:v="urn:schemas-microsoft-com:vml"', out)
        self.assertIn('mc:Ignorable="p14 v"', out)
        self.assertIn('<mc:Choice Requires="p14">', out)
        self.assertIn('<a:pPr eaLnBrk="0" latinLnBrk="0"/>', out)
        self.assertNotIn("ns0", out)

    # 3. 카드 기본 스타일이 흰색 글자/그림자를 강제하지 않아야 함
    def test_card_removes_default_white_text_style(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        card = create_korean_card(slide, Inches(1), Inches(1), Inches(3), Inches(2))
        self.assertIsNone(card._element.find(qn("p:style")))

        dark = create_korean_card(
            slide, Inches(1), Inches(1), Inches(3), Inches(2),
            bg_color=RGBColor(0x11, 0x18, 0x27), text_color=RGBColor(0xFF, 0xFF, 0xFF)
        )
        clr = dark.text_frame._txBody.find(
            "./a:lstStyle/a:lvl1pPr/a:defRPr/a:solidFill/a:srgbClr".replace("a:", "{%s}" % A_NS)
        )
        self.assertIsNotNone(clr)
        self.assertEqual(clr.get("val"), "FFFFFF")

    # 4. 기존 buClr/buFont/tabLst가 있어도 스키마 순서와 choice 단일성 유지
    def test_set_native_bullet_schema_order_with_existing_elements(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        p = tb.text_frame.paragraphs[0]
        p.text = "불릿"
        pPr = p._p.get_or_add_pPr()
        for frag in [
            '<a:spcBef><a:spcPts val="600"/></a:spcBef>',
            '<a:buClr><a:srgbClr val="FF0000"/></a:buClr>',
            '<a:buFont typeface="Arial"/>',
            '<a:tabLst/>',
            '<a:defRPr/>',
        ]:
            pPr.append(parse_xml('<x %s>%s</x>' % (nsdecls("a"), frag))[0])

        set_native_bullet(p)

        tags = [c.tag.split('}')[-1] for c in pPr]
        self.assertEqual(tags, ["spcBef", "buClrTx", "buSzTx", "buChar", "tabLst", "defRPr"])

    # 5. latin 테마 참조/한글 미지원 폰트를 ea에 그대로 복사하지 않아야 함
    def test_resolve_ea_typeface(self):
        self.assertEqual(resolve_ea_typeface("+mn-lt", "Pretendard"), "+mn-ea")
        self.assertEqual(resolve_ea_typeface("+mj-lt", "Pretendard"), "+mj-ea")
        self.assertEqual(resolve_ea_typeface("Calibri", "Pretendard"), "Pretendard")
        self.assertEqual(resolve_ea_typeface(None, "Pretendard"), "Pretendard")
        self.assertEqual(resolve_ea_typeface("Noto Sans KR", "Pretendard"), "Noto Sans KR")
        self.assertEqual(resolve_ea_typeface("맑은 고딕", "Pretendard"), "맑은 고딕")
        self.assertEqual(resolve_ea_typeface("NanumGothic", "Pretendard"), "NanumGothic")

    def test_patch_maps_theme_font_and_fixes_bad_ea(self):
        xml_str = """<a:txBody xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<a:p><a:r><a:rPr><a:latin typeface="+mj-lt"/></a:rPr><a:t>제목</a:t></a:r></a:p>
<a:p><a:r><a:rPr><a:latin typeface="+mn-lt"/><a:ea typeface="+mn-lt"/></a:rPr><a:t>본문</a:t></a:r></a:p>
</a:txBody>"""
        _, new_xml, _ = patch_xml_content(xml_str.encode("utf-8"))
        root = ET.fromstring(new_xml)
        eas = [e.get("typeface") for e in root.iter(qualify(A_NS, "ea"))]
        self.assertEqual(eas, ["+mj-ea", "+mn-ea"])

    # 6. Table 객체를 직접 전달해도 적용되어야 함
    def test_apply_korean_formatting_table_object(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        gf = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
        gf.table.cell(0, 0).text = "표 셀"

        apply_korean_formatting(gf.table, font_name="Pretendard")

        for cell in gf.table.iter_cells():
            pPr = cell.text_frame.paragraphs[0]._p.pPr
            self.assertIsNotNone(pPr)
            self.assertEqual(pPr.get("eaLnBrk"), "0")


if __name__ == '__main__':
    unittest.main()
