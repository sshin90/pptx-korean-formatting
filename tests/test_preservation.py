"""Behavioral regressions for the 2026-10 review; fixtures stay in TemporaryDirectory."""
import contextlib
import io
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml import parse_xml

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from scripts.fix_korean_pptx import (
    A_NS, P_NS, patch_xml_content, patch_theme_xml, verify_xml_content,
    process_pptx, get_slide_file_map,
)
from scripts.korean_pptx import apply_korean_formatting, set_korean_font_and_lang

NS = {"a": A_NS, "p": P_NS}
ROOT = Path(__file__).resolve().parents[1]


def body(content, wrap="square"):
    return (f'<p:txBody xmlns:p="{P_NS}" xmlns:a="{A_NS}">'
            f'<a:bodyPr wrap="{wrap}"/>{content}</p:txBody>').encode()


def paragraph(text="한글", prop="", rprop="", kind="r"):
    extra = ' id="{00000000-0000-0000-0000-000000000001}" type="datetime1"' if kind == "fld" else ""
    return f'<a:p>{prop}<a:{kind}{extra}>{rprop}<a:t>{text}</a:t></a:{kind}></a:p>'


DEFAULT = '<a:pPr eaLnBrk="0" latinLnBrk="0"><a:defRPr lang="ko-KR"><a:latin typeface="NanumGothic"/><a:ea typeface="NanumGothic"/></a:defRPr></a:pPr>'


def patched(data, **kw):
    changed, output, stats = patch_xml_content(data, **kw)
    if "error" in stats:
        raise AssertionError(stats)
    return E.fromstring(output)


def read_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


class FormattingRegressions(unittest.TestCase):
    def test_preserve_paragraph_inheritance_after_parent_change(self):
        root = patched(body(paragraph(prop=DEFAULT)))
        for name in ("latin", "ea"):
            root.find(f".//a:defRPr/a:{name}", NS).set("typeface", "Noto Sans KR")
        self.assertFalse(root.xpath(".//a:rPr/a:ea | .//a:rPr/a:latin", namespaces=NS))
        # The verifier must now see the parent's new font, not an old copied value.
        self.assertEqual(verify_xml_content(E.tostring(root), expected_font="Noto Sans KR"), {})

    def test_preserve_list_style_inheritance_after_parent_change(self):
        style = '<a:lstStyle><a:lvl1pPr eaLnBrk="false" latinLnBrk="false"><a:defRPr lang="ko-KR"><a:latin typeface="NanumGothic"/><a:ea typeface="NanumGothic"/></a:defRPr></a:lvl1pPr></a:lstStyle>'
        root = patched(body(style + paragraph()))
        for face in root.xpath("./a:lstStyle//a:latin | ./a:lstStyle//a:ea", namespaces=NS):
            face.set("typeface", "Noto Sans KR")
        self.assertFalse(root.xpath("./a:p//a:ea | ./a:p//a:latin", namespaces=NS))
        self.assertEqual(verify_xml_content(E.tostring(root), expected_font="Noto Sans KR"), {})

    def test_higher_inheritance_reported_unknown_without_injecting_font(self):
        root = patched(body(paragraph()))
        self.assertFalse(root.xpath(".//a:ea | .//a:latin", namespaces=NS))
        self.assertGreater(verify_xml_content(E.tostring(root))["unknown_ea_font_inheritance"], 0)

    def test_preserve_custom_fonts_and_intentional_pairing(self):
        data = body(paragraph(prop=DEFAULT, rprop='<a:rPr><a:latin typeface="Arial"/><a:ea typeface="CompanyCustomFont"/></a:rPr>'))
        root = patched(data)
        self.assertEqual(root.find(".//a:rPr/a:latin", NS).get("typeface"), "Arial")
        self.assertEqual(root.find(".//a:rPr/a:ea", NS).get("typeface"), "CompanyCustomFont")
        self.assertEqual(verify_xml_content(E.tostring(root)), {})

    def test_preserve_english_only_bytes(self):
        data = body(paragraph("English only", rprop='<a:rPr lang="en-US"><a:latin typeface="Arial"/></a:rPr>'), wrap="none")
        changed, output, _ = patch_xml_content(data, default_font="NanumGothic", force_font=True)
        self.assertFalse(changed)
        self.assertEqual(output, data)

    def test_mixed_runs_keep_foreign_language_and_explicit_font(self):
        data = body('<a:p><a:r><a:rPr lang="en-US"><a:latin typeface="Arial"/></a:rPr><a:t>English</a:t></a:r><a:r><a:t>한글</a:t></a:r></a:p>')
        root = patched(data, default_font="NanumGothic", force_font=True)
        eng, ko = root.findall(".//a:r", NS)
        self.assertEqual(eng.find("a:rPr", NS).get("lang"), "en-US")
        self.assertEqual(eng.find("a:rPr/a:latin", NS).get("typeface"), "Arial")
        self.assertEqual(ko.find("a:rPr", NS).get("lang"), "ko-KR")

    def test_force_font_removes_theme_override_and_preserves_emphasis_link(self):
        data = body(paragraph(rprop='<a:rPr b="1"><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill><a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:hlinkClick xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId2"/></a:rPr>'))
        root = patched(data, default_font="NanumGothic", force_font=True)
        prop = root.find(".//a:rPr", NS)
        self.assertEqual(prop.get("b"), "1")
        self.assertIsNotNone(prop.find("a:hlinkClick", NS))
        self.assertEqual(prop.find("a:solidFill/a:srgbClr", NS).get("val"), "FF0000")
        self.assertIsNone(prop.find("a:latin", NS))
        self.assertIsNone(prop.find("a:ea", NS))
        self.assertEqual(verify_xml_content(E.tostring(root), expected_font="NanumGothic"), {})

    def test_direct_symbol_run_retains_font(self):
        data = body(paragraph("A", rprop='<a:rPr lang="en-US"><a:latin typeface="Wingdings"/></a:rPr>'))
        root = patched(data, default_font="NanumGothic", force_font=True, all_text=True)
        self.assertEqual(root.find(".//a:rPr/a:latin", NS).get("typeface"), "Wingdings")

    def test_wrap_disabled_is_detected_and_repaired(self):
        data = body(paragraph(prop=DEFAULT), wrap="none")
        self.assertEqual(verify_xml_content(data)["word_wrap_disabled"], 1)
        root = patched(data)
        self.assertEqual(root.find("a:bodyPr", NS).get("wrap"), "square")
        self.assertEqual(verify_xml_content(E.tostring(root)), {})

    def test_direct_script_import_without_package(self):
        script = 'import sys; sys.path.insert(0, r"' + str(ROOT / 'scripts') + '"); from korean_pptx import apply_korean_formatting; from pptx import Presentation; apply_korean_formatting(Presentation())'
        res = subprocess.run([sys.executable, "-c", script], capture_output=True, text=True)
        self.assertEqual(res.returncode, 0, res.stderr)

    def test_korean_field_language_detected_and_fixed(self):
        data = body(paragraph("2026년", prop=DEFAULT, rprop='<a:rPr lang="en-US"/>', kind="fld"))
        self.assertEqual(verify_xml_content(data)["non_korean_language"], 1)
        self.assertEqual(verify_xml_content(E.tostring(patched(data))), {})

    def test_numeric_field_uses_all_text_opt_in(self):
        data = body(paragraph("1", prop=DEFAULT, rprop='<a:rPr lang="en-US"/>', kind="fld"))
        self.assertEqual(verify_xml_content(data), {})
        self.assertEqual(verify_xml_content(data, all_text=True)["non_korean_language"], 1)
        root = patched(data, all_text=True)
        self.assertEqual(verify_xml_content(E.tostring(root), all_text=True), {})

    def test_explicit_empty_ea_not_hidden_by_good_parent(self):
        data = body(paragraph(prop=DEFAULT, rprop='<a:rPr><a:ea typeface=""/></a:rPr>'))
        self.assertGreater(verify_xml_content(data)["invalid_ea_typeface"], 0)

    def test_explicit_unexpected_font_not_hidden_by_good_parent(self):
        data = body(paragraph(prop=DEFAULT, rprop='<a:rPr><a:ea typeface="Calibri"/></a:rPr>'))
        self.assertGreater(verify_xml_content(data, expected_font="NanumGothic")["unexpected_ea_font"], 0)

    def test_boolean_false_and_local_list_inheritance_accepted(self):
        data = body(paragraph(prop=DEFAULT.replace('="0"', '="false"')))
        self.assertEqual(verify_xml_content(data), {})

    def test_active_list_style_break_is_checked(self):
        data = body('<a:lstStyle><a:lvl2pPr eaLnBrk="1" latinLnBrk="0"><a:defRPr lang="ko-KR"><a:ea typeface="NanumGothic"/></a:defRPr></a:lvl2pPr></a:lstStyle>' + paragraph(prop='<a:pPr lvl="1"/>'))
        self.assertEqual(verify_xml_content(data)["enabled_eaLnBrk"], 1)
        self.assertEqual(verify_xml_content(E.tostring(patched(data))), {})

    def test_shared_styles_checked_in_all_text_mode(self):
        data = body('<a:lstStyle><a:lvl1pPr eaLnBrk="1" latinLnBrk="0"/></a:lstStyle>')
        self.assertEqual(verify_xml_content(data, all_text=True)["style_eaLnBrk"], 1)
        self.assertEqual(verify_xml_content(E.tostring(patched(data, all_text=True)), all_text=True), {})

    def test_schema_order_and_choice_violations_reported(self):
        data = body(paragraph(prop='<a:pPr><a:defRPr/><a:buNone/><a:buChar char="•"/></a:pPr>',
                              rprop='<a:rPr><a:hlinkClick/><a:ea typeface="NanumGothic"/></a:rPr>'))
        issues = verify_xml_content(data)
        self.assertEqual(issues["paragraph_element_order_or_choice"], 1)
        self.assertEqual(issues["font_element_order"], 1)

    def test_default_and_force_patch_are_idempotent(self):
        for kw in ({}, {"default_font":"NanumGothic", "force_font":True}, {"all_text":True}):
            with self.subTest(kw=kw):
                _, first, _ = patch_xml_content(body(paragraph()), **kw)
                changed, second, _ = patch_xml_content(first, **kw)
                self.assertFalse(changed)
                self.assertEqual(first, second)

    def test_default_theme_is_byte_identical(self):
        prs = Presentation()
        with tempfile.TemporaryDirectory() as td:
            file = Path(td) / "theme.pptx"
            prs.save(file)
            original = read_parts(file)["ppt/theme/theme1.xml"]
        changed, output, _ = patch_theme_xml(original)
        self.assertFalse(changed)
        self.assertEqual(original, output)

    def test_helper_default_preserves_font_and_normalize_keeps_parent_editable(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        shape.text_frame.text = "한글"
        p = shape.text_frame.paragraphs[0]
        set_korean_font_and_lang(p, "NanumGothic")
        apply_korean_formatting(prs)
        self.assertIsNone(p.runs[0].font.name)
        apply_korean_formatting(prs, font_name="Noto Sans KR")
        set_korean_font_and_lang(p, "Malgun Gothic")
        self.assertIsNone(p.runs[0].font.name)
        self.assertEqual(p.font.name, "Malgun Gothic")

    def test_group_and_table_helper_use_same_policy(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        shape.text_frame.text = "그룹 텍스트"
        apply_korean_formatting(prs, font_name="NanumGothic")
        self.assertEqual(shape.text_frame.paragraphs[0].font.name, "NanumGothic")
        self.assertIsNone(shape.text_frame.paragraphs[0].runs[0].font.name)


class PackageRegressions(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "sample.pptx"
        prs = Presentation()
        for text in ("첫 번째", "두 번째"):
            sl = prs.slides.add_slide(prs.slide_layouts[0])
            sl.shapes.title.text = text
        prs.save(self.path)

    def process(self, **kw):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return process_pptx(self.path, **kw)

    def rewrite(self, edits):
        parts = read_parts(self.path)
        parts.update(edits)
        with zipfile.ZipFile(self.path, "w") as z:
            for name, content in parts.items():
                z.writestr(name, content)

    def test_one_slide_never_changes_shared_parts(self):
        before = read_parts(self.path)
        self.assertEqual(self.process(target_slide=1, default_font="NanumGothic", force_font=True), 0)
        after = read_parts(self.path)
        self.assertEqual({n for n in before if before[n] != after[n]}, {"ppt/slides/slide1.xml"})
        self.assertEqual(verify_xml_content(after["ppt/slides/slide1.xml"], expected_font="NanumGothic"), {})

    def test_no_masters_does_not_prevent_explicit_font(self):
        self.assertEqual(self.process(target_slide=1, include_masters=False, default_font="NanumGothic", force_font=True), 0)
        self.assertEqual(self.process(target_slide=1, verify_only=True, default_font="NanumGothic", force_font=True), 0)

    def test_default_preserves_all_themes_masters_and_layouts_without_korean(self):
        before = read_parts(self.path)
        self.assertEqual(self.process(), 0)
        after = read_parts(self.path)
        for name in before:
            if any(p in name for p in ("/theme/", "/slideMasters/", "/slideLayouts/")):
                self.assertEqual(before[name], after[name], name)

    def test_reordered_slides_use_presentation_order(self):
        parts = read_parts(self.path)
        root = E.fromstring(parts["ppt/presentation.xml"])
        ids = root.find("p:sldIdLst", NS)
        ids.insert(0, ids[-1])
        self.rewrite({"ppt/presentation.xml":E.tostring(root)})
        before = read_parts(self.path)
        self.assertEqual(self.process(target_slide=1), 0)
        after = read_parts(self.path)
        self.assertEqual({n for n in before if before[n] != after[n]}, {"ppt/slides/slide2.xml"})

    def test_invalid_slide_scope_fails_without_changes(self):
        before = self.path.read_bytes()
        for number in (0, -1, 3):
            self.assertEqual(self.process(target_slide=number), 1)
            self.assertEqual(self.path.read_bytes(), before)

    def test_shared_theme_requires_explicit_full_document_option(self):
        before = self.path.read_bytes()
        for kw in ({"update_theme":True}, {"update_theme":True,"target_slide":1,"default_font":"NanumGothic","force_font":True},
                   {"update_theme":True,"include_masters":False,"default_font":"NanumGothic","force_font":True}):
            self.assertEqual(self.process(**kw), 1)
            self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(self.process(default_font="NanumGothic", force_font=True, update_theme=True), 0)
        root = E.fromstring(read_parts(self.path)["ppt/theme/theme1.xml"])
        self.assertEqual(root.xpath(".//a:font[@script='Hang']/@typeface", namespaces=NS), ["NanumGothic","NanumGothic"])

    def test_corrupt_xml_preserves_source_and_existing_output(self):
        self.rewrite({"ppt/slides/slide1.xml":b"<broken"})
        before = self.path.read_bytes()
        out = Path(self.temp.name) / "output.pptx"
        out.write_bytes(b"existing-output")
        self.assertEqual(self.process(output_path=str(out)), 1)
        self.assertEqual(self.process(verify_only=True), 1)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(out.read_bytes(), b"existing-output")
        self.assertFalse(list(Path(self.temp.name).glob(".korean_pptx_*")))

    def test_corrupt_relationships_fail_instead_of_filename_fallback(self):
        self.rewrite({"ppt/_rels/presentation.xml.rels":b"<broken"})
        before = self.path.read_bytes()
        self.assertEqual(self.process(target_slide=1), 1)
        self.assertEqual(before, self.path.read_bytes())

    def test_replace_failure_preserves_source_and_cleans_temp(self):
        before = self.path.read_bytes()
        with patch("scripts.fix_korean_pptx.os.replace", side_effect=PermissionError("locked")):
            self.assertEqual(self.process(), 1)
        self.assertEqual(before, self.path.read_bytes())
        self.assertFalse(list(Path(self.temp.name).glob(".korean_pptx_*")))

    def test_verify_exit_codes_and_read_only(self):
        before = self.path.read_bytes()
        self.assertEqual(self.process(verify_only=True), 3)  # inherited placeholder defaults unresolved
        self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(self.process(default_font="NanumGothic", force_font=True), 0)
        after = self.path.read_bytes()
        self.assertEqual(self.process(verify_only=True), 0)
        self.assertEqual(after, self.path.read_bytes())
        self.assertEqual(self.process(verify_only=True, default_font="OtherFont", force_font=True), 2)

    def test_output_copy_leaves_original_unchanged_and_reopens(self):
        before = self.path.read_bytes()
        out = Path(self.temp.name) / "copy.pptx"
        self.assertEqual(self.process(output_path=str(out)), 0)
        self.assertEqual(before, self.path.read_bytes())
        self.assertEqual(len(Presentation(out).slides), 2)

    def test_cli_invalid_combinations_return_error(self):
        result = subprocess.run([sys.executable, "-X", "utf8", "-B", str(ROOT/"scripts/fix_korean_pptx.py"),
                                 str(self.path), "--slide", "1", "-f", "NanumGothic", "--update-theme"],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 1)
        self.assertIn("--update-theme", result.stderr)


if __name__ == "__main__":
    unittest.main()
