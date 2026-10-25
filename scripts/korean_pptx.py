#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
korean_pptx.py
python-pptx 기반 프레젠테이션 제작 시 사용하는 한국어 전용 서식 헬퍼 모듈.
"""

from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

try:
    from .fix_korean_pptx import HANGUL_RE, patch_tree
except ImportError:
    from fix_korean_pptx import HANGUL_RE, patch_tree

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
POST_LATIN_TAGS = [
    f"{{{A_NS}}}{t}"
    for t in ["ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]
]
POST_EA_TAGS = [
    f"{{{A_NS}}}{t}"
    for t in ["cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]
]

# pPr 내 불릿 관련 요소 전체 (색상/크기/폰트/종류 각 그룹은 하나만 존재 가능한 choice)
BULLET_TAGS = [
    f"{{{A_NS}}}{t}"
    for t in [
        "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont",
        "buNone", "buAutoNum", "buChar", "buBlip",
    ]
]
# 스키마상 불릿 요소 뒤에 와야 하는 pPr 자식 요소
POST_BULLET_TAGS = [f"{{{A_NS}}}{t}" for t in ["tabLst", "defRPr", "extLst"]]


def _insert_latin_to_prop(prop_elm, font_name):
    """prop_elm (rPr 또는 defRPr) 내부에 OpenXML 스키마 순서를 준수하여 <a:latin> 요소를 삽입합니다."""
    latin = prop_elm.find(f"{{{A_NS}}}latin")
    if latin is not None:
        latin.set("typeface", font_name)
        return

    latin_elem = OxmlElement("a:latin")
    latin_elem.set("typeface", font_name)

    for i, child in enumerate(prop_elm):
        if child.tag in POST_LATIN_TAGS:
            prop_elm.insert(i, latin_elem)
            return

    prop_elm.append(latin_elem)


def _insert_ea_to_prop(prop_elm, font_name):
    """prop_elm (rPr 또는 defRPr) 내부에 OpenXML 스키마 순서를 준수하여 <a:ea> 요소를 삽입합니다."""
    ea = prop_elm.find(f"{{{A_NS}}}ea")
    if ea is not None:
        ea.set("typeface", font_name)
        return

    ea_elem = OxmlElement("a:ea")
    ea_elem.set("typeface", font_name)

    latin = prop_elm.find(f"{{{A_NS}}}latin")
    if latin is not None:
        idx = list(prop_elm).index(latin)
        prop_elm.insert(idx + 1, ea_elem)
        return

    for i, child in enumerate(prop_elm):
        if child.tag in POST_EA_TAGS:
            prop_elm.insert(i, ea_elem)
            return

    prop_elm.append(ea_elem)


def set_korean_font_and_lang(target, font_name="Pretendard"):
    """
    Run, Paragraph 또는 Font 객체에 언어(ko-KR)와 폰트(latin + ea: East Asian)를 동시 설정.
    명시적으로 지정하는 저수준 함수입니다. 기존 파일의 상속 보존에는
    apply_korean_formatting(target)을 사용하세요. 폰트 설치/글리프는 별도 확인합니다.
    """
    font_obj = target if hasattr(target, "language_id") else getattr(target, "font", None)
    if font_obj is not None:
        try:
            font_obj.language_id = MSO_LANGUAGE_ID.KOREAN
        except Exception:
            pass
        font_obj.name = font_name

    # 1. Run 객체인 경우
    if hasattr(target, "_r"):
        prop_elm = target._r.get_or_add_rPr()
        _insert_latin_to_prop(prop_elm, font_name)
        _insert_ea_to_prop(prop_elm, font_name)
        return

    # 2. Paragraph 객체인 경우 (문단 레벨 defRPr 상속 설정 + 기존 런 언어 태그 동기화)
    if hasattr(target, "_p"):
        p_elm = target._p
        prop_elm = p_elm.get_or_add_pPr().get_or_add_defRPr()
        prop_elm.set("lang", "ko-KR")
        _insert_latin_to_prop(prop_elm, font_name)
        _insert_ea_to_prop(prop_elm, font_name)

        # 혼합 언어 문단에서 외국어 run의 명시적 언어는 유지합니다.
        for r in getattr(target, "runs", ()):
            if HANGUL_RE.search(r.text):
                r.font.language_id = MSO_LANGUAGE_ID.KOREAN

        end_para = p_elm.find(f"{{{A_NS}}}endParaRPr")
        if end_para is not None:
            end_para.set("lang", "ko-KR")
        return

    # 3. Font 객체를 직접 전달한 경우 (target._element가 rPr 또는 defRPr)
    prop_elm = getattr(target, "_element", None)
    if prop_elm is not None:
        prop_elm.set("lang", "ko-KR")
        _insert_latin_to_prop(prop_elm, font_name)
        _insert_ea_to_prop(prop_elm, font_name)


def set_word_break_prevention(paragraph):
    """
    문단에 단어 중간 줄바꿈 방지 속성 (eaLnBrk="0", latinLnBrk="0") 설정.
    한글/영문 단어가 음절 단위로 쪼개지는 현상을 방지합니다.
    """
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("eaLnBrk", "0")
    pPr.set("latinLnBrk", "0")


def set_native_bullet(paragraph, char="•", mar_l=Inches(0.3), indent=-Inches(0.2)):
    """
    OpenXML ISO/IEC 29500-1 스키마 순서를 엄격 준수하여 네이티브 불릿을 적용.
    buChar가 tabLst/defRPr/extLst보다 앞에 오도록 보장하여 PowerPoint 파일 복구 경고를 방지합니다.
    """
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(mar_l)))
    pPr.set("indent", str(int(indent)))

    # 1. 기존 불릿 관련 태그 전체 제거 (buClr/buSzPct/buFont 등과 choice 중복 방지)
    for child in list(pPr):
        if child.tag in BULLET_TAGS:
            pPr.remove(child)

    # 2. 불릿 요소 생성: 텍스트 색상 상속(buClrTx), 크기 상속(buSzTx), 불릿 문자(buChar)
    bu_clr_tx = OxmlElement("a:buClrTx")
    bu_sz_tx = OxmlElement("a:buSzTx")
    bu_char = OxmlElement("a:buChar")
    bu_char.set("char", char)

    # 3. 스키마 순서 준수: lnSpc/spcBef/spcAft -> [buClrTx, buSzTx, buChar] -> tabLst -> defRPr -> extLst
    idx = len(pPr)
    for i, child in enumerate(pPr):
        if child.tag in POST_BULLET_TAGS:
            idx = i
            break
    pPr.insert(idx, bu_clr_tx)
    pPr.insert(idx + 1, bu_sz_tx)
    pPr.insert(idx + 2, bu_char)


def create_korean_card(
    slide,
    left,
    top,
    width,
    height,
    bg_color=RGBColor(245, 247, 250),
    border_color=RGBColor(220, 225, 230),
    line_width_pt=1,
    margin_inch=0.3,
    shape_type=MSO_SHAPE.RECTANGLE,
    vertical_anchor=None,
    text_color=None,
):
    """
    배경 도형과 텍스트 프레임을 일체화한 단일 카드/도형(사각형, 원형, 둥근 사각형 등) 생성.
    사용자 드래그 시 분리되지 않고 일체형으로 편집할 수 있습니다.

    :param slide: 슬라이드 객체
    :param left: 가로 위치 (Inches 또는 Pt)
    :param top: 세로 위치 (Inches 또는 Pt)
    :param width: 너비 (Inches 또는 Pt)
    :param height: 높이 (Inches 또는 Pt)
    :param bg_color: 배경 색상 (RGBColor 또는 None=투명)
    :param border_color: 테두리 색상 (RGBColor 또는 None=없음)
    :param line_width_pt: 테두리 두께 (숫자, pt 단위, 기본값 1)
    :param margin_inch: 텍스트 내부 마진 (숫자, inch 단위, 기본값 0.3)
    :param shape_type: MSO_SHAPE enum (기본값: RECTANGLE, 원형은 OVAL, 둥근 사각형은 ROUNDED_RECTANGLE 등)
    :param vertical_anchor: MSO_ANCHOR enum (미지정 시 OVAL은 MSO_ANCHOR.MIDDLE, 그 외는 MSO_ANCHOR.TOP 기본 적용)
    :param text_color: 기본 글자색 (RGBColor, 미지정 시 테마 텍스트 색 상속 / 어두운 배경이면 지정)
    """
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)

    # python-pptx 기본 도형 스타일(p:style)은 흰색 글자(fontRef lt1)와 그림자(effectRef)를 강제해
    # 밝은 배경 카드에서 글자가 보이지 않으므로 제거하고 테마 기본 서식을 상속받게 함
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)

    if bg_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
    else:
        shape.fill.background()

    if border_color is not None and line_width_pt > 0:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(line_width_pt)
    else:
        shape.line.fill.background()

    tf = shape.text_frame
    tf.word_wrap = True

    if vertical_anchor is not None:
        tf.vertical_anchor = vertical_anchor
    else:
        # 원형/타원형(OVAL)의 경우 통상 중앙 정렬, 직사각형 카드는 상단 정렬이 기본
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if shape_type == MSO_SHAPE.OVAL else MSO_ANCHOR.TOP

    if margin_inch is not None:
        tf.margin_left = Inches(margin_inch)
        tf.margin_right = Inches(margin_inch)
        tf.margin_top = Inches(margin_inch)
        tf.margin_bottom = Inches(margin_inch)

    if text_color is not None:
        # run마다 하드코딩하지 않고 도형의 목록 스타일(lstStyle)에서 1~9단계 모두 상속
        lst_style = tf._txBody.find(qn("a:lstStyle"))
        if lst_style is not None:
            for lvl in range(1, 10):
                lvl_tag = f"lvl{lvl}pPr"
                existing = lst_style.find(qn(f"a:{lvl_tag}"))
                if existing is not None:
                    lst_style.remove(existing)
                lst_style.append(
                    parse_xml(
                        f'<a:{lvl_tag} {nsdecls("a")} eaLnBrk="0" latinLnBrk="0">'
                        f'<a:defRPr lang="ko-KR"><a:solidFill>'
                        f'<a:srgbClr val="{text_color}"/></a:solidFill></a:defRPr>'
                        f'</a:{lvl_tag}>'
                    )
                )

    if tf.paragraphs:
        set_word_break_prevention(tf.paragraphs[0])

    return shape


# 단일 도형 생성 헬퍼 별칭 (원형, 뱃지 등 카드 외 다양한 도형 생성 시 직관적 사용)
create_korean_shape = create_korean_card


def _extract_text_frames(container):
    """
    Presentation, Slide, Shape, Table, GroupShape 등 다양한 객체에서
    모든 TextFrame을 재귀적으로 수집합니다.
    """
    tfs = []
    # 1. Presentation 객체 (slides 컬렉션 순회)
    if hasattr(container, "slides"):
        for slide in container.slides:
            tfs.extend(_extract_text_frames(slide))
        return tfs

    # 2. Table 객체를 직접 전달한 경우 (graphicFrame.table)
    if hasattr(container, "iter_cells"):
        return [cell.text_frame for cell in container.iter_cells()]

    # 3. Slide 또는 GroupShape의 shapes 컬렉션
    shapes = []
    if hasattr(container, "shapes"):
        shapes = list(container.shapes)
    elif isinstance(container, (list, tuple)):
        shapes = list(container)
    else:
        shapes = [container]

    for shape in shapes:
        # 이미 TextFrame인 경우
        if hasattr(shape, "paragraphs") and hasattr(shape, "word_wrap"):
            tfs.append(shape)
            continue

        # GroupShape: 내부에 shapes 컬렉션이 있음
        if hasattr(shape, "shapes"):
            tfs.extend(_extract_text_frames(shape))
            continue

        # Table 객체
        if getattr(shape, "has_table", False):
            for cell in shape.table.iter_cells():
                if getattr(cell, "text_frame", None):
                    tfs.append(cell.text_frame)
            continue

        # 일반 도형의 TextFrame
        if getattr(shape, "has_text_frame", False) and shape.text_frame:
            tfs.append(shape.text_frame)

    return tfs


def apply_korean_formatting(target_obj, font_name=None, *, all_text=False):
    """
    프레젠테이션(Presentation), 슬라이드(Slide), 도형(Shape), 표(Table),
    또는 TextFrame의 한국어 문단을 보정합니다. 기본값은 기존 폰트/상속 보존.
    font_name을 명시하면 문단 기본 폰트를 지정하고 한국어 run의 폰트 override만
    제거합니다. 외국어/빈 문단도 처리하려면 all_text=True를 사용합니다.
    표/그룹을 순회하며 공유 마스터·레이아웃·테마는 수정하지 않습니다.
    """
    if font_name is not None and (not font_name.strip() or font_name.startswith("+")):
        raise ValueError("실제 폰트 이름을 지정하세요")
    for tf in _extract_text_frames(target_obj):
        patch_tree(tf._txBody, font_name=font_name, all_text=all_text)
